"""Tests for clean DOCX export (interactive OCR)."""
import io
from pathlib import Path

import pytest
from PIL import Image
from docx import Document

from app.ocr_assist import job_store
from app.ocr_assist.docx_export import TIBETAN_FONT, build_clean_docx, write_clean_docx
from app.ocr_assist.job_store import OUTPUT_DOCX_FILE, create_job, finalize_page, load_job


@pytest.fixture
def fake_pdf(monkeypatch):
    def _fake_render(pdf_bytes: bytes, *, dpi: int):
        return [Image.new("RGB", (40, 40), "white") for _ in range(2)]

    monkeypatch.setattr(job_store, "_render_pdf", _fake_render)
    return b"fake-pdf-bytes"


def test_build_clean_docx_page_headings_and_font():
    payload = build_clean_docx([(1, "བོད་ཡིག"), (2, "line a\nline b")])
    doc = Document(io.BytesIO(payload))
    texts = [p.text for p in doc.paragraphs]
    assert "Page 1" in texts
    assert "Page 2" in texts
    assert "བོད་ཡིག" in texts
    assert "line a" in texts
    assert "line b" in texts

    body_runs = [
        run
        for p in doc.paragraphs
        if not p.style.name.startswith("Heading")
        for run in p.runs
        if run.text
    ]
    assert body_runs
    assert all(run.font.name == TIBETAN_FONT for run in body_runs)


def test_write_clean_docx_none_when_empty(tmp_path: Path, fake_pdf):
    job = create_job(
        fake_pdf,
        source_file="s.pdf",
        baseline_settings={"k_factor": 2.0},
        jobs_root=tmp_path,
    )
    assert write_clean_docx(job) is None
    assert not (job.root / OUTPUT_DOCX_FILE).is_file()


def test_write_clean_docx_dedupes_on_re_finalize(tmp_path: Path, fake_pdf):
    job = create_job(
        fake_pdf,
        source_file="s.pdf",
        baseline_settings={"k_factor": 2.0},
        jobs_root=tmp_path,
    )
    finalize_page(job, 1, final_text="first")
    finalize_page(job, 1, final_text="updated")
    job = load_job(tmp_path, job.id)
    doc = Document(str(job.root / OUTPUT_DOCX_FILE))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings == ["Page 1"]
    assert "updated" in [p.text for p in doc.paragraphs]
